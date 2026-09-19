import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_master_template():
    os.makedirs('templates', exist_ok=True)
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Premium Color Palette
    DARK_BG = RGBColor(17, 17, 17)        # #111111 Ink Black
    CREAM_CARD = RGBColor(245, 242, 235)  # #F5F2EB Tactile Cream
    PURE_WHITE = RGBColor(255, 255, 255)
    RED_PRIMARY = RGBColor(196, 30, 36)   # #C41E24 Crimson Red
    YELLOW_ACCENT = RGBColor(255, 232, 0) # #FFE800 Canary Yellow
    GRAY_TEXT = RGBColor(100, 100, 100)
    
    def add_background(slide, color=DARK_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        
        # Outer structural border
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
        frame.fill.background()
        frame.line.color.rgb = PURE_WHITE
        frame.line.width = Pt(1.5)
        return bg

    def add_header(slide, title_text, step_tag="01 / BLUEPRINT"):
        # Category Tag Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(3.2), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RED_PRIMARY
        badge.line.color.rgb = PURE_WHITE
        badge.line.width = Pt(1.5)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = f"FRAMEFLOW • {step_tag}"
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = PURE_WHITE
        p_b.font.name = "Space Grotesk"
        p_b.alignment = PP_ALIGN.CENTER
        
        # Slide Main Title
        title_box = slide.shapes.add_textbox(Inches(4.2), Inches(0.45), Inches(8.3), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = PURE_WHITE
        p.font.name = "Syne"

    # ==================== SLIDE 1: COVER ====================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1, DARK_BG)
    
    # Large Hero Frame Card
    card1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CREAM_CARD
    card1.line.color.rgb = DARK_BG
    card1.line.width = Pt(3)
    
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.8)
    tf1.margin_top = Inches(0.8)
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "⚡ CREATIVE FRAMEWORK BLUEPRINT"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = RED_PRIMARY
    p_tag.font.name = "Space Grotesk"
    
    p_brand = tf1.add_paragraph()
    p_brand.text = "{{BRAND_NAME}}"
    p_brand.font.size = Pt(64)
    p_brand.font.bold = True
    p_brand.font.color.rgb = RED_PRIMARY
    p_brand.font.name = "Syne"
    
    p_proj = tf1.add_paragraph()
    p_proj.text = "{{PROJECT_NAME}}"
    p_proj.font.size = Pt(32)
    p_proj.font.bold = True
    p_proj.font.color.rgb = DARK_BG
    p_proj.font.name = "Space Grotesk"

    p_div = tf1.add_paragraph()
    p_div.text = "──────────────────────────────────────────────"
    p_div.font.size = Pt(14)
    p_div.font.color.rgb = DARK_BG

    p_yr = tf1.add_paragraph()
    p_yr.text = "CAMPAIGN YEAR: {{YEAR}}   •   AGENCY CONFIDENTIAL"
    p_yr.font.size = Pt(14)
    p_yr.font.bold = True
    p_yr.font.color.rgb = DARK_BG
    p_yr.font.name = "Space Grotesk"

    # ==================== SLIDE 2: COLOUR PALETTE ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2, DARK_BG)
    add_header(slide2, "COLOUR PALETTE DIRECTION", "01 / COLOURS")
    
    # Swatch 1 - Primary
    sw1 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(3.64), Inches(4.2))
    sw1.name = "SWATCH_PRIMARY"
    sw1.fill.solid()
    sw1.fill.fore_color.rgb = DARK_BG
    sw1.line.color.rgb = PURE_WHITE
    sw1.line.width = Pt(2.5)
    tf_sw1 = sw1.text_frame
    p = tf_sw1.paragraphs[0]
    p.text = "PRIMARY COLOR\n\n{{PRIMARY_COLOR}}"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = PURE_WHITE
    p.font.name = "Space Grotesk"
    p.alignment = PP_ALIGN.CENTER
    
    # Swatch 2 - Secondary
    sw2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.84), Inches(1.6), Inches(3.64), Inches(4.2))
    sw2.name = "SWATCH_SECONDARY"
    sw2.fill.solid()
    sw2.fill.fore_color.rgb = RED_PRIMARY
    sw2.line.color.rgb = PURE_WHITE
    sw2.line.width = Pt(2.5)
    tf_sw2 = sw2.text_frame
    p = tf_sw2.paragraphs[0]
    p.text = "SECONDARY COLOR\n\n{{SECONDARY_COLOR}}"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = PURE_WHITE
    p.font.name = "Space Grotesk"
    p.alignment = PP_ALIGN.CENTER

    # Swatch 3 - Accent
    sw3 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.88), Inches(1.6), Inches(3.64), Inches(4.2))
    sw3.name = "SWATCH_ACCENT"
    sw3.fill.solid()
    sw3.fill.fore_color.rgb = CREAM_CARD
    sw3.line.color.rgb = PURE_WHITE
    sw3.line.width = Pt(2.5)
    tf_sw3 = sw3.text_frame
    p = tf_sw3.paragraphs[0]
    p.text = "ACCENT COLOR\n\n{{ACCENT_COLOR}}"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_BG
    p.font.name = "Space Grotesk"
    p.alignment = PP_ALIGN.CENTER

    # Palette Guidelines Footer Box
    guide_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.72), Inches(0.8))
    guide_box.fill.solid()
    guide_box.fill.fore_color.rgb = CREAM_CARD
    guide_box.line.color.rgb = DARK_BG
    guide_box.line.width = Pt(2)
    tf_g = guide_box.text_frame
    p_g = tf_g.paragraphs[0]
    p_g.text = "PALETTE SPECS: Apply primary for structural frames & headlines; secondary for campaign CTAs; accent for high-visibility highlights."
    p_g.font.size = Pt(11)
    p_g.font.bold = True
    p_g.font.color.rgb = DARK_BG
    p_g.font.name = "Space Grotesk"
    p_g.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 3: TYPOGRAPHY ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3, DARK_BG)
    add_header(slide3, "TYPOGRAPHY HIERARCHY & TYPEFACES", "02 / TYPOGRAPHY")
    
    # Display Font Box
    box_disp = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.72), Inches(1.6))
    box_disp.fill.solid()
    box_disp.fill.fore_color.rgb = CREAM_CARD
    box_disp.line.color.rgb = DARK_BG
    box_disp.line.width = Pt(2)
    tf_d = box_disp.text_frame
    tf_d.margin_left = Inches(0.4)
    p_d0 = tf_d.paragraphs[0]
    p_d0.text = "DISPLAY FONT: {{DISPLAY_FONT}}"
    p_d0.font.size = Pt(14)
    p_d0.font.bold = True
    p_d0.font.color.rgb = RED_PRIMARY
    p_d0.font.name = "Space Grotesk"
    
    p_d1 = tf_d.add_paragraph()
    p_d1.text = "The Quick Brown Fox Jumps Over The Lazy Dog — 0123456789"
    p_d1.font.size = Pt(24)
    p_d1.font.bold = True
    p_d1.font.color.rgb = DARK_BG
    p_d1.font.name = "Cormorant Garamond"
    
    # Secondary Font Box
    box_sec = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.4), Inches(11.72), Inches(1.6))
    box_sec.fill.solid()
    box_sec.fill.fore_color.rgb = CREAM_CARD
    box_sec.line.color.rgb = DARK_BG
    box_sec.line.width = Pt(2)
    tf_s = box_sec.text_frame
    tf_s.margin_left = Inches(0.4)
    p_s0 = tf_s.paragraphs[0]
    p_s0.text = "SECONDARY FONT: {{SECONDARY_FONT}}"
    p_s0.font.size = Pt(14)
    p_s0.font.bold = True
    p_s0.font.color.rgb = RED_PRIMARY
    p_s0.font.name = "Space Grotesk"
    
    p_s1 = tf_s.add_paragraph()
    p_s1.text = "Editorial story outlines, moodboard captions & shoot direction notes."
    p_s1.font.size = Pt(20)
    p_s1.font.bold = True
    p_s1.font.color.rgb = DARK_BG
    p_s1.font.name = "Playfair Display"

    # Caption Font Box
    box_cap = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.72), Inches(1.6))
    box_cap.fill.solid()
    box_cap.fill.fore_color.rgb = CREAM_CARD
    box_cap.line.color.rgb = DARK_BG
    box_cap.line.width = Pt(2)
    tf_c = box_cap.text_frame
    tf_c.margin_left = Inches(0.4)
    p_c0 = tf_c.paragraphs[0]
    p_c0.text = "CAPTION FONT: {{CAPTION_FONT}}"
    p_c0.font.size = Pt(14)
    p_c0.font.bold = True
    p_c0.font.color.rgb = RED_PRIMARY
    p_c0.font.name = "Space Grotesk"
    
    p_c1 = tf_c.add_paragraph()
    p_c1.text = "TECHNICAL SPECS: 4K 60FPS | COLOR LOG-C | ASPECT RATIO 9:16 | BITRATE 100Mbps"
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = DARK_BG
    p_c1.font.name = "Inter"

    # ==================== SLIDE 4: EDITING STYLE ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4, DARK_BG)
    add_header(slide4, "EDITING STYLE & PACING CONTROL", "03 / EDITING")
    
    # Pacing Container Card
    card_pace = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.72), Inches(2.4))
    card_pace.fill.solid()
    card_pace.fill.fore_color.rgb = CREAM_CARD
    card_pace.line.color.rgb = DARK_BG
    card_pace.line.width = Pt(2.5)
    tf_p = card_pace.text_frame
    tf_p.margin_left = Inches(0.4)
    p_p0 = tf_p.paragraphs[0]
    p_p0.text = "EDIT PACING INDICATOR"
    p_p0.font.size = Pt(16)
    p_p0.font.bold = True
    p_p0.font.color.rgb = RED_PRIMARY
    p_p0.font.name = "Space Grotesk"
    
    # Pacing Track Shape
    track = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.7), Inches(10.9), Inches(0.4))
    track.name = "PACING_TRACK"
    track.fill.solid()
    track.fill.fore_color.rgb = PURE_WHITE
    track.line.color.rgb = DARK_BG
    track.line.width = Pt(2)
    
    # Pacing Marker Shape
    marker = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.5), Inches(1.4), Inches(0.8))
    marker.name = "PACING_MARKER"
    marker.fill.solid()
    marker.fill.fore_color.rgb = RED_PRIMARY
    marker.line.color.rgb = DARK_BG
    marker.line.width = Pt(2)
    tf_m = marker.text_frame
    p_m = tf_m.paragraphs[0]
    p_m.text = "{{PACING}}"
    p_m.font.size = Pt(16)
    p_m.font.bold = True
    p_m.font.color.rgb = PURE_WHITE
    p_m.font.name = "Space Grotesk"
    p_m.alignment = PP_ALIGN.CENTER

    # Edit Vibe Tags Box
    box_tags = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.72), Inches(2.6))
    box_tags.fill.solid()
    box_tags.fill.fore_color.rgb = CREAM_CARD
    box_tags.line.color.rgb = DARK_BG
    box_tags.line.width = Pt(2.5)
    tf_t = box_tags.text_frame
    tf_t.margin_left = Inches(0.4)
    tf_t.margin_top = Inches(0.3)
    p_t0 = tf_t.paragraphs[0]
    p_t0.text = "EDIT VIBE & TRANSITION STYLES:"
    p_t0.font.size = Pt(16)
    p_t0.font.bold = True
    p_t0.font.color.rgb = RED_PRIMARY
    p_t0.font.name = "Space Grotesk"
    
    p_t1 = tf_t.add_paragraph()
    p_t1.text = "• {{EDIT_STYLE_1}}\n• {{EDIT_STYLE_2}}\n• {{EDIT_STYLE_3}}"
    p_t1.font.size = Pt(20)
    p_t1.font.bold = True
    p_t1.font.color.rgb = DARK_BG
    p_t1.font.name = "Space Grotesk"

    # ==================== SLIDE 5: WARDROBE ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5, DARK_BG)
    add_header(slide5, "WARDROBE DIRECTION & MOODBOARD", "04 / WARDROBE")
    
    # Text metadata card
    card_w = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    card_w.fill.solid()
    card_w.fill.fore_color.rgb = CREAM_CARD
    card_w.line.color.rgb = DARK_BG
    card_w.line.width = Pt(2.5)
    tf_w = card_w.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = Inches(0.4)
    tf_w.margin_top = Inches(0.4)
    
    p_w0 = tf_w.paragraphs[0]
    p_w0.text = "STYLE TONE"
    p_w0.font.size = Pt(14)
    p_w0.font.bold = True
    p_w0.font.color.rgb = RED_PRIMARY
    p_w0.font.name = "Space Grotesk"
    
    p_w1 = tf_w.add_paragraph()
    p_w1.text = "{{WARDROBE_TONE}}\n"
    p_w1.font.size = Pt(20)
    p_w1.font.bold = True
    p_w1.font.color.rgb = DARK_BG
    p_w1.font.name = "Space Grotesk"

    p_w2 = tf_w.add_paragraph()
    p_w2.text = "COLOUR COMBINATIONS"
    p_w2.font.size = Pt(14)
    p_w2.font.bold = True
    p_w2.font.color.rgb = RED_PRIMARY
    p_w2.font.name = "Space Grotesk"
    
    p_w3 = tf_w.add_paragraph()
    p_w3.text = "{{WARDROBE_COLORS}}"
    p_w3.font.size = Pt(18)
    p_w3.font.bold = True
    p_w3.font.color.rgb = DARK_BG
    p_w3.font.name = "Space Grotesk"

    # Placeholder image box 1
    img1_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.6), Inches(1.6), Inches(3.4), Inches(5.2))
    img1_box.name = "{{WARDROBE_IMAGE_1}}"
    img1_box.fill.solid()
    img1_box.fill.fore_color.rgb = DARK_BG
    img1_box.line.color.rgb = PURE_WHITE
    img1_box.line.width = Pt(2)
    tf_i1 = img1_box.text_frame
    p_i1 = tf_i1.paragraphs[0]
    p_i1.text = "[ WARDROBE REF 01 ]"
    p_i1.font.size = Pt(14)
    p_i1.font.color.rgb = CREAM_CARD
    p_i1.alignment = PP_ALIGN.CENTER

    # Placeholder image box 2
    img2_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(1.6), Inches(3.32), Inches(5.2))
    img2_box.name = "{{WARDROBE_IMAGE_2}}"
    img2_box.fill.solid()
    img2_box.fill.fore_color.rgb = DARK_BG
    img2_box.line.color.rgb = PURE_WHITE
    img2_box.line.width = Pt(2)
    tf_i2 = img2_box.text_frame
    p_i2 = tf_i2.paragraphs[0]
    p_i2.text = "[ WARDROBE REF 02 ]"
    p_i2.font.size = Pt(14)
    p_i2.font.color.rgb = CREAM_CARD
    p_i2.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 6: BACKGROUND ====================
    slide6 = prs.slides.add_slide(blank_layout)
    add_background(slide6, DARK_BG)
    add_header(slide6, "BACKGROUND & SHOOT LOCATIONS", "05 / BACKGROUND")
    
    # Text metadata card
    card_bg = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    card_bg.fill.solid()
    card_bg.fill.fore_color.rgb = CREAM_CARD
    card_bg.line.color.rgb = DARK_BG
    card_bg.line.width = Pt(2.5)
    tf_bg = card_bg.text_frame
    tf_bg.word_wrap = True
    tf_bg.margin_left = Inches(0.4)
    tf_bg.margin_top = Inches(0.4)
    
    p_bg0 = tf_bg.paragraphs[0]
    p_bg0.text = "BACKGROUND TONE"
    p_bg0.font.size = Pt(14)
    p_bg0.font.bold = True
    p_bg0.font.color.rgb = RED_PRIMARY
    p_bg0.font.name = "Space Grotesk"
    
    p_bg1 = tf_bg.add_paragraph()
    p_bg1.text = "{{BACKGROUND_TONE}}\n"
    p_bg1.font.size = Pt(20)
    p_bg1.font.bold = True
    p_bg1.font.color.rgb = DARK_BG
    p_bg1.font.name = "Space Grotesk"

    p_bg2 = tf_bg.add_paragraph()
    p_bg2.text = "KEY LOCATIONS"
    p_bg2.font.size = Pt(14)
    p_bg2.font.bold = True
    p_bg2.font.color.rgb = RED_PRIMARY
    p_bg2.font.name = "Space Grotesk"
    
    p_bg3 = tf_bg.add_paragraph()
    p_bg3.text = "{{BACKGROUND_LOCATIONS}}"
    p_bg3.font.size = Pt(18)
    p_bg3.font.bold = True
    p_bg3.font.color.rgb = DARK_BG
    p_bg3.font.name = "Space Grotesk"

    # Placeholder background image box 1
    bg_img1 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.6), Inches(1.6), Inches(3.4), Inches(5.2))
    bg_img1.name = "{{BACKGROUND_IMAGE_1}}"
    bg_img1.fill.solid()
    bg_img1.fill.fore_color.rgb = DARK_BG
    bg_img1.line.color.rgb = PURE_WHITE
    bg_img1.line.width = Pt(2)
    tf_b1 = bg_img1.text_frame
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "[ BACKGROUND REF 01 ]"
    p_b1.font.size = Pt(14)
    p_b1.font.color.rgb = CREAM_CARD
    p_b1.alignment = PP_ALIGN.CENTER

    # Placeholder background image box 2
    bg_img2 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(1.6), Inches(3.32), Inches(5.2))
    bg_img2.name = "{{BACKGROUND_IMAGE_2}}"
    bg_img2.fill.solid()
    bg_img2.fill.fore_color.rgb = DARK_BG
    bg_img2.line.color.rgb = PURE_WHITE
    bg_img2.line.width = Pt(2)
    tf_b2 = bg_img2.text_frame
    p_b2 = tf_b2.paragraphs[0]
    p_b2.text = "[ BACKGROUND REF 02 ]"
    p_b2.font.size = Pt(14)
    p_b2.font.color.rgb = CREAM_CARD
    p_b2.alignment = PP_ALIGN.CENTER

    output_path = os.path.join('templates', 'amore_framework.pptx')
    prs.save(output_path)
    print(f"Ultra-premium master template created successfully at {output_path}")

if __name__ == '__main__':
    create_master_template()
